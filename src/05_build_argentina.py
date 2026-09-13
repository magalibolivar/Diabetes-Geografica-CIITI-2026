# -*- coding: utf-8 -*-
"""
05 - Construye el dataset provincial de Argentina (24 jurisdicciones) combinando:
  - ENFR 2018 (diabetes y determinantes de estilo de vida, por provincia)
  - NBI 2010 (INDEC, serie censal historica) - socioeconomico
  - Pobreza EPH 2do semestre 2018 (INDEC, por aglomerado -> promediado por provincia)
  - Cobertura de salud, Censo 2022 (INDEC) - acceso a salud
  - Densidad poblacional, Censo 2022 (INDEC) - proxy de ruralidad

Fuentes esperadas en data/ (ver handoff.md para detalle y links):
  - cuadros_definitivos_enfr_2018.xls
  - serie_nbi_2022.xlsx
  - cuadros_informe_pobreza_03_26.xls
  - c2022_tp_salud_c1.xlsx

Salida: data/processed/enfr2018_provincias.csv (24 filas).

IMPORTANTE (ver handoff.md, seccion "Resultados del modelado"): con N=24 el
modelo multivariado tipo BRFSS NO generaliza (LOOCV R2 ~0). Este script solo
arma el dataset; no asume que vaya a usarse para un modelo predictivo tal
cual - la decision de enfoque (atlas descriptivo / paper metodologico /
validacion autorreporte vs. medicion / esperar 5ta ENFR) queda pendiente.
"""
import xlrd
import openpyxl
import pandas as pd
import numpy as np
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
ENFR = DATA / "cuadros_definitivos_enfr_2018.xls"
NBI = DATA / "serie_nbi_2022.xlsx"
POBREZA = DATA / "cuadros_informe_pobreza_03_26.xls"
SALUD = DATA / "c2022_tp_salud_c1.xlsx"
OUT = DATA / "processed" / "enfr2018_provincias.csv"

# Orden fijo de provincias tal como aparecen en los Cuadros de la ENFR (filas 7-34,
# con subtotales regionales intercalados que se omiten).
PROVINCIAS_ENFR = [
    "CABA", "Buenos Aires", "Córdoba", "Entre Ríos", "La Pampa", "Santa Fe",
    "Catamarca", "Jujuy", "La Rioja", "Salta", "Santiago del Estero", "Tucumán",
    "Corrientes", "Chaco", "Formosa", "Misiones",
    "Chubut", "Neuquén", "Río Negro", "Santa Cruz", "Tierra del Fuego",
    "Mendoza", "San Juan", "San Luis",
]


def enfr_cuadro(sheet_name, colname):
    wb = xlrd.open_workbook(ENFR)
    s = wb.sheet_by_name(sheet_name)
    vals = {}
    for r in range(s.nrows):
        label = str(s.cell_value(r, 0)).replace("\xa0", " ").strip()
        if label in PROVINCIAS_ENFR:
            vals[label] = s.cell_value(r, 1)
    faltan = set(PROVINCIAS_ENFR) - set(vals.keys())
    if faltan:
        raise ValueError(f"{sheet_name}: faltan provincias {faltan}")
    return pd.Series(vals, name=colname)


def build_enfr():
    cols = {
        "tasa_diabetes_pct": "Cuadro 7.3",       # glucemia elevada/diabetes, autorreporte
        "obesidad_pct": "Cuadro 6.3",
        "inactividad_fisica_pct": "Cuadro 3.1",
        "tabaquismo_pct": "Cuadro 2.1",
        "presion_elevada_pct": "Cuadro 8.3",
        "colesterol_elevado_pct": "Cuadro 9.3",
        "consumo_frutas_verduras_pct": "Cuadro 5.7",
        "tratamiento_diabetes_pct": "Cuadro 7.5",  # de los diagnosticados, % en tratamiento
        "medicion_glucemia_pct": "Cuadro 7.1",     # % que se midio la glucemia alguna vez
    }
    df = pd.DataFrame({name: enfr_cuadro(sheet, name) for name, sheet in cols.items()})
    df.index.name = "provincia"
    return df


def build_nbi():
    wb = openpyxl.load_workbook(NBI, data_only=True)
    ws = wb["NBI_Hogares_%_80_22"]
    header = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    col_2010 = header.index(2010) + 1
    vals = {}
    for r in range(6, ws.max_row + 1):
        name = ws.cell(r, 1).value
        if not name:
            continue
        name = str(name).strip()
        v = ws.cell(r, col_2010).value
        if name == "Ciudad Autónoma de Buenos Aires":
            name = "CABA"
        elif name.startswith("Tierra del Fuego"):
            name = "Tierra del Fuego"
        elif name == "Entre Ríos ":
            name = "Entre Ríos"
        if name in PROVINCIAS_ENFR and v is not None:
            vals[name] = float(v)
    faltan = set(PROVINCIAS_ENFR) - set(vals.keys())
    if faltan:
        raise ValueError(f"NBI: faltan provincias {faltan}")
    return pd.Series(vals, name="nbi_2010_pct")


# Mapeo aglomerado EPH -> provincia (Cuadro 4.3, columna "2 semestre 2018").
# Cuando una provincia tiene mas de un aglomerado se promedia sin ponderar
# (no hay poblacion por aglomerado en este cuadro para ponderar correctamente).
AGLOMERADO_PROVINCIA = {
    "Ciudad Autónoma de Buenos Aires": "CABA",
    "Partidos del GBA": "Buenos Aires",
    "Bahía Blanca - Cerri": "Buenos Aires",
    "Mar del Plata ": "Buenos Aires",
    "Mar del Plata": "Buenos Aires",
    "San Nicolás - Villa Constitución": "Buenos Aires",
    "Gran La Plata": "Buenos Aires",
    "Gran Mendoza": "Mendoza",
    "Gran San Juan": "San Juan",
    "Gran San Luis": "San Luis",
    "Gran Resistencia (2)": "Chaco",
    "Corrientes (4)": "Corrientes",
    "Formosa": "Formosa",
    "Posadas": "Misiones",
    "Gran Catamarca": "Catamarca",
    "Gran Tucumán - Tafí Viejo": "Tucumán",
    "Jujuy - Palpalá": "Jujuy",
    "La Rioja": "La Rioja",
    "Salta": "Salta",
    "Santiago del Estero - La Banda": "Santiago del Estero",
    "Concordia": "Entre Ríos",
    "Gran Paraná": "Entre Ríos",
    "Gran Córdoba": "Córdoba",
    "Río Cuarto": "Córdoba",
    "Gran Rosario": "Santa Fe",
    "Gran Santa Fe": "Santa Fe",
    "Santa Rosa - Toay": "La Pampa",
    "Comodoro Rivadavia - Rada Tilly": "Chubut",
    "Rawson - Trelew": "Chubut",
    "Neuquen - Plottier": "Neuquén",
    "Río Gallegos": "Santa Cruz",
    "Ushuaia - Río Grande (3)": "Tierra del Fuego",
    "Viedma - Carmen de Patagones": "Río Negro",
}


def build_pobreza():
    wb = xlrd.open_workbook(POBREZA)
    s = wb.sheet_by_name("Cuadro 4.3")
    header = [s.cell_value(2, c) for c in range(s.ncols)]
    col_personas = None
    for i, v in enumerate(header):
        if v == "2 semestre 2018":
            col_personas = i + 2  # Hogares en i, blank en i+1, Personas en i+2
            break
    assert col_personas is not None, "No se encontro la columna '2 semestre 2018'"
    rows = {}
    for r in range(5, s.nrows):
        label = str(s.cell_value(r, 0)).strip()
        if label in AGLOMERADO_PROVINCIA:
            val = s.cell_value(r, col_personas)
            if isinstance(val, str):
                val = float(val.replace(",", "."))
            rows.setdefault(AGLOMERADO_PROVINCIA[label], []).append(float(val))
    vals = {prov: float(np.mean(v)) for prov, v in rows.items()}
    faltan = set(PROVINCIAS_ENFR) - set(vals.keys())
    if faltan:
        raise ValueError(f"Pobreza EPH: faltan provincias {faltan}")
    return pd.Series(vals, name="pobreza_personas_pct")


def build_cobertura():
    wb = openpyxl.load_workbook(SALUD, data_only=True)
    ws = wb["Cobertura de salud N°1"]
    vals = {}
    for r in range(6, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if not name:
            continue
        name = str(name).strip()
        if name.startswith(("24 Partidos", "Resto de partidos", "31 Partidos")):
            continue  # subtotales de partidos de Buenos Aires: se usa la fila "Buenos Aires"
        total = ws.cell(r, 3).value
        sin_cob = ws.cell(r, 6).value
        if total in (None, 0) or sin_cob is None:
            continue
        if name == "Ciudad Autónoma de Buenos Aires":
            name = "CABA"
        elif name.startswith("Tierra del Fuego"):
            name = "Tierra del Fuego"
        if name in PROVINCIAS_ENFR:
            vals[name] = round(100 * float(sin_cob) / float(total), 2)
    faltan = set(PROVINCIAS_ENFR) - set(vals.keys())
    if faltan:
        raise ValueError(f"Cobertura de salud: faltan provincias {faltan}")
    return pd.Series(vals, name="sin_cobertura_salud_pct")


# Densidad poblacional 2022 (INDEC, Censo 2022 - resultados definitivos), hab/km2.
# Fuente secundaria (ver handoff.md); Tierra del Fuego ajustada a superficie
# efectiva (~21.263 km2, sin el reclamo antartico) en vez de la superficie
# jurisdiccional completa (~910.324 km2) que distorsiona la densidad oficial.
DENSIDAD_2022 = {
    "CABA": 15169.0, "Buenos Aires": 57.27, "Catamarca": 4.23, "Chaco": 11.32,
    "Chubut": 2.64, "Córdoba": 23.30, "Corrientes": 13.61, "Entre Ríos": 18.21,
    "Formosa": 8.05, "Jujuy": 15.26, "La Pampa": 2.52, "La Rioja": 4.20,
    "Mendoza": 13.71, "Misiones": 42.70, "Neuquén": 7.52, "Río Negro": 3.71,
    "Salta": 9.29, "San Juan": 9.31, "San Luis": 7.19, "Santa Cruz": 1.38,
    "Santa Fe": 26.60, "Santiago del Estero": 7.75,
    "Tierra del Fuego": 8.6,
    "Tucumán": 76.65,
}
REGION_ENFR = {
    "CABA": "Pampeana y GBA", "Buenos Aires": "Pampeana y GBA", "Córdoba": "Pampeana y GBA",
    "Entre Ríos": "Pampeana y GBA", "La Pampa": "Pampeana y GBA", "Santa Fe": "Pampeana y GBA",
    "Catamarca": "Noroeste", "Jujuy": "Noroeste", "La Rioja": "Noroeste", "Salta": "Noroeste",
    "Santiago del Estero": "Noroeste", "Tucumán": "Noroeste",
    "Corrientes": "Noreste", "Chaco": "Noreste", "Formosa": "Noreste", "Misiones": "Noreste",
    "Chubut": "Patagonia", "Neuquén": "Patagonia", "Río Negro": "Patagonia",
    "Santa Cruz": "Patagonia", "Tierra del Fuego": "Patagonia",
    "Mendoza": "Cuyo", "San Juan": "Cuyo", "San Luis": "Cuyo",
}

# Coeficientes de variacion (CV, %) de la prevalencia de diabetes por provincia
# (Cuadro 7.3) - miden la confiabilidad de cada estimacion muestral. Todas
# estan en un rango razonable (5.3%-12.5%); ninguna provincia requiere
# exclusion por baja confiabilidad. Utiles para una WLS de robustez.
CV_DIABETES = {
    "CABA": 10.2, "Buenos Aires": 5.3, "Córdoba": 8.3, "Entre Ríos": 6.9,
    "La Pampa": 12.0, "Santa Fe": 8.2, "Catamarca": 11.6, "Jujuy": 11.7,
    "La Rioja": 12.3, "Salta": 9.2, "Santiago del Estero": 12.5, "Tucumán": 9.5,
    "Corrientes": 9.8, "Chaco": 11.2, "Formosa": 10.3, "Misiones": 9.6,
    "Chubut": 11.2, "Neuquén": 9.9, "Río Negro": 9.4, "Santa Cruz": 9.2,
    "Tierra del Fuego": 12.2, "Mendoza": 10.8, "San Juan": 9.2, "San Luis": 9.8,
}


def main():
    df = build_enfr()
    df = df.join(build_nbi())
    df = df.join(build_pobreza())
    df = df.join(build_cobertura())
    df["densidad_hab_km2"] = pd.Series(DENSIDAD_2022)
    df["log_densidad"] = np.log10(df["densidad_hab_km2"])
    df["region"] = pd.Series(REGION_ENFR)
    df["cv_diabetes_pct"] = pd.Series(CV_DIABETES)
    df = df.reset_index().rename(columns={"index": "provincia"})
    df = df.sort_values("provincia").reset_index(drop=True)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False)
    print(df.to_string())
    print(f"\nOK -> {OUT}  ({len(df)} jurisdicciones)")


if __name__ == "__main__":
    main()
